#!/usr/bin/env python3
"""
Generate the Biz-Salama Seller Pitch Deck (.pptx).

Audience: Tanzanian sellers (mama biashara, hawkers, suppliers, small-shop owners).
Primary language: Kiswahili. English subtitles on every slide.

Output: /app/biz_salama_seller_pitch_sw.pptx
Serve: GET /api/docs/seller-pitch.pptx (registered in server.py)
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor as DMLColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pathlib import Path

# Brand palette (matches Tailwind config + LandingPage.tsx)
GOLD       = DMLColor(0xF5, 0x9E, 0x0B)   # gold-500
GOLD_LIGHT = DMLColor(0xFB, 0xBF, 0x24)   # gold-400
INK_900    = DMLColor(0x0F, 0x17, 0x2A)
INK_800    = DMLColor(0x1E, 0x29, 0x3B)
INK_700    = DMLColor(0x33, 0x41, 0x55)
INK_400    = DMLColor(0x94, 0xA3, 0xB8)
INK_300    = DMLColor(0xCB, 0xD5, 0xE1)
EMERALD    = DMLColor(0x10, 0xB9, 0x81)
EMERALD_LT = DMLColor(0x34, 0xD3, 0x99)
ROSE       = DMLColor(0xF4, 0x3F, 0x5E)
WHITE      = DMLColor(0xFF, 0xFF, 0xFF)


# ─── Helpers ───────────────────────────────────────────────────────────────

def fill(shape, rgb: DMLColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    shape.line.fill.background()


def add_rect(slide, x, y, w, h, rgb: DMLColor, *, rounded=False):
    s = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(s, x, y, w, h)
    fill(shp, rgb)
    shp.shadow.inherit = False
    return shp


def add_text(
    slide, x, y, w, h, text, *,
    size=18, bold=False, italic=False, color=WHITE, align=PP_ALIGN.LEFT, font="Calibri",
    anchor=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor

    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = font
    return box


def slide_bg(slide, rgb: DMLColor):
    """Paint the whole slide background by adding a fill rectangle."""
    return add_rect(slide, 0, 0, prs.slide_width, prs.slide_height, rgb)


def add_pill(slide, x, y, w, h, label, *, fill_rgb=GOLD, text_rgb=INK_900, size=10):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = 0.5  # full pill
    fill(shp, fill_rgb)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Emu(100000)
    tf.margin_top = tf.margin_bottom = Emu(40000)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = text_rgb
    r.font.name = "Calibri"
    return shp


def footer(slide, idx, total):
    """Bottom corner footer present on every content slide."""
    add_text(
        slide, Inches(0.5), Inches(7.0), Inches(6), Inches(0.3),
        "Biz-Salama · biz-salama.co.tz · Uza salama. Lipwa salama.",
        size=9, color=INK_400,
    )
    add_text(
        slide, Inches(9.4), Inches(7.0), Inches(0.6), Inches(0.3),
        f"{idx}/{total}",
        size=9, color=INK_400, align=PP_ALIGN.RIGHT,
    )


# ─── Build deck ────────────────────────────────────────────────────────────

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

TOTAL_SLIDES = 11

# ── Slide 1 — Cover ─────────────────────────────────────────────────────
s1 = prs.slides.add_slide(blank)
slide_bg(s1, INK_900)
# Decorative gold blob
shp = s1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2), Inches(-2), Inches(5), Inches(5))
fill(shp, GOLD)
shp.fill.transparency = 0.85  # softened
# Brand mark
add_pill(s1, Inches(0.6), Inches(0.6), Inches(2.5), Inches(0.45),
         "🛡  BIZ-SALAMA", fill_rgb=GOLD, text_rgb=INK_900, size=11)
add_text(s1, Inches(0.6), Inches(2.0), Inches(9), Inches(1.6),
         "Karibu Biz-Salama",
         size=64, bold=True, color=WHITE, font="Calibri")
add_text(s1, Inches(0.6), Inches(3.4), Inches(9), Inches(0.7),
         "Soko salama la Tanzania — uza, lipwa, kua.",
         size=24, color=GOLD_LIGHT, font="Calibri")
add_text(s1, Inches(0.6), Inches(4.2), Inches(9), Inches(0.5),
         "Tanzania's trusted escrow marketplace — sell, get paid, grow.",
         size=14, color=INK_300, font="Calibri")
# Three stat pills at bottom
add_pill(s1, Inches(0.6), Inches(5.5), Inches(2.3), Inches(0.6),
         "✅ Malipo Yamelindwa  /  Escrow Protected",
         fill_rgb=EMERALD, text_rgb=WHITE, size=10)
add_pill(s1, Inches(3.1), Inches(5.5), Inches(2.3), Inches(0.6),
         "📱 M-Pesa · Tigo · Airtel",
         fill_rgb=INK_700, text_rgb=WHITE, size=10)
add_pill(s1, Inches(5.6), Inches(5.5), Inches(2.3), Inches(0.6),
         "🎙  Orodhesha kwa Sauti  /  Voice Listing",
         fill_rgb=INK_700, text_rgb=WHITE, size=10)
add_text(s1, Inches(0.6), Inches(6.8), Inches(9), Inches(0.4),
         "Mwongozo wa Mfanyabiashara  ·  Seller Onboarding Pitch Deck",
         size=11, color=INK_400, font="Calibri")


# ── Slide 2 — Tatizo (Problem) ─────────────────────────────────────────
s2 = prs.slides.add_slide(blank)
slide_bg(s2, INK_900)
add_pill(s2, Inches(0.6), Inches(0.5), Inches(1.2), Inches(0.4),
         "TATIZO  /  THE PROBLEM", fill_rgb=ROSE, text_rgb=WHITE, size=9)
add_text(s2, Inches(0.6), Inches(1.0), Inches(9), Inches(1.0),
         "Kuuza mtandaoni TZ ni hatari",
         size=40, bold=True, color=WHITE)
add_text(s2, Inches(0.6), Inches(1.95), Inches(9), Inches(0.5),
         "Selling online in Tanzania is risky — here's why:",
         size=14, color=INK_400)

# Four pain cards (2×2 grid)
pains = [
    ("💸", "Mteja anataka uone kwanza", "Buyer wants you to ship before paying",
     "Wewe unalipa usafirishaji, halafu mteja anatoweka."),
    ("🎭", "Watu wa uongo (scammers)", "Fake buyers / fake sellers",
     "Mtu anajifanya mteja, halafu anaiba bidhaa au pesa."),
    ("⏳", "Hakuna ulinzi wa malipo", "No payment protection",
     "Hakuna anayeshikilia pesa salama kati ya nyote wawili."),
    ("⚖️", "Hakuna njia ya kutatua mgogoro", "No dispute resolution",
     "Kama mambo yanakwenda mrama, hakuna mahali pa kupiga simu."),
]
for i, (emoji, sw, en, body) in enumerate(pains):
    col, row = i % 2, i // 2
    x = Inches(0.6 + col * 4.6)
    y = Inches(2.8 + row * 1.85)
    card = add_rect(s2, x, y, Inches(4.3), Inches(1.65), INK_800, rounded=True)
    card.adjustments[0] = 0.12
    add_text(s2, x + Inches(0.3), y + Inches(0.15), Inches(0.6), Inches(0.5),
             emoji, size=24)
    add_text(s2, x + Inches(1.0), y + Inches(0.18), Inches(3.1), Inches(0.4),
             sw, size=13, bold=True, color=WHITE)
    add_text(s2, x + Inches(1.0), y + Inches(0.58), Inches(3.1), Inches(0.3),
             en, size=9, color=GOLD_LIGHT)
    add_text(s2, x + Inches(0.3), y + Inches(0.95), Inches(3.8), Inches(0.6),
             body, size=10, color=INK_300)
footer(s2, 2, TOTAL_SLIDES)


# ── Slide 3 — Suluhisho (Solution) ──────────────────────────────────────
s3 = prs.slides.add_slide(blank)
slide_bg(s3, INK_900)
add_pill(s3, Inches(0.6), Inches(0.5), Inches(1.6), Inches(0.4),
         "SULUHISHO  /  THE SOLUTION", fill_rgb=EMERALD, text_rgb=WHITE, size=9)
add_text(s3, Inches(0.6), Inches(1.0), Inches(9), Inches(1.0),
         "Biz-Salama = Escrow + Soko",
         size=42, bold=True, color=WHITE)
add_text(s3, Inches(0.6), Inches(1.95), Inches(9), Inches(0.5),
         "Biz-Salama = Escrow + Marketplace, built for Tanzanian commerce",
         size=14, color=INK_400)

# Three big benefit blocks
benefits = [
    ("🔒", "Pesa zinashikiliwa salama", "Funds held in escrow",
     "Mteja anaweka pesa Biz-Salama. Hatupeleki kwako mpaka mteja athibitishe bidhaa imefika salama."),
    ("📲", "Lipwa kwa M-Pesa moja kwa moja", "Direct M-Pesa payout",
     "Tunatuma kwa M-Pesa, Tigo Pesa au Airtel Money mara tu mteja anapokubali."),
    ("🛡", "Ulinzi wa pande zote mbili", "Both sides protected",
     "Wewe huitaji kufuatilia mteja. Mteja huitaji kukutuhumu. Sote tunashinda."),
]
for i, (emoji, sw, en, body) in enumerate(benefits):
    y = Inches(2.7 + i * 1.45)
    card = add_rect(s3, Inches(0.6), y, Inches(8.8), Inches(1.25), INK_800, rounded=True)
    card.adjustments[0] = 0.1
    # Icon bubble
    bubble = s3.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.85), y + Inches(0.22),
                                  Inches(0.8), Inches(0.8))
    fill(bubble, GOLD)
    add_text(s3, Inches(0.85), y + Inches(0.32), Inches(0.8), Inches(0.6),
             emoji, size=24, align=PP_ALIGN.CENTER)
    add_text(s3, Inches(1.95), y + Inches(0.2), Inches(7.0), Inches(0.4),
             sw, size=16, bold=True, color=WHITE)
    add_text(s3, Inches(1.95), y + Inches(0.55), Inches(7.0), Inches(0.3),
             en, size=10, color=GOLD_LIGHT)
    add_text(s3, Inches(1.95), y + Inches(0.85), Inches(7.0), Inches(0.4),
             body, size=10, color=INK_300)
footer(s3, 3, TOTAL_SLIDES)


# ── Slide 4 — Jinsi inavyofanya kazi (How it works) ────────────────────
s4 = prs.slides.add_slide(blank)
slide_bg(s4, INK_900)
add_pill(s4, Inches(0.6), Inches(0.5), Inches(1.6), Inches(0.4),
         "JINSI INAVYOFANYA KAZI  /  HOW IT WORKS", fill_rgb=GOLD, text_rgb=INK_900, size=9)
add_text(s4, Inches(0.6), Inches(1.0), Inches(9), Inches(1.0),
         "Hatua 5 tu  /  Just 5 steps",
         size=38, bold=True, color=WHITE)

steps = [
    ("1", "📝", "Orodhesha bidhaa", "List your product", "Andika au sema kwa sauti (Kiswahili / English)."),
    ("2", "💰", "Mteja analipa", "Buyer pays escrow", "M-Pesa / Tigo / Airtel — pesa zinashikiliwa Biz-Salama."),
    ("3", "📦", "Tuma bidhaa", "You ship the item", "Tunakuarifu pesa imeshikiliwa — pakizia na peleka."),
    ("4", "✅", "Mteja anathibitisha", "Buyer confirms", "Anaweka 'Nimepokea'. Tunapata uthibitisho."),
    ("5", "🏦", "Lipwa moja kwa moja", "Get paid instantly", "Pesa zinaingia M-Pesa yako (98% ya bei yako)."),
]
for i, (n, emoji, sw, en, desc) in enumerate(steps):
    x = Inches(0.4 + i * 1.92)
    y = Inches(2.4)
    # Card
    card = add_rect(s4, x, y, Inches(1.7), Inches(3.4), INK_800, rounded=True)
    card.adjustments[0] = 0.08
    # Number pill
    num = s4.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.6), y - Inches(0.3),
                               Inches(0.6), Inches(0.6))
    fill(num, GOLD)
    add_text(s4, x + Inches(0.6), y - Inches(0.22), Inches(0.6), Inches(0.5),
             n, size=18, bold=True, color=INK_900, align=PP_ALIGN.CENTER)
    add_text(s4, x, y + Inches(0.5), Inches(1.7), Inches(0.6),
             emoji, size=28, align=PP_ALIGN.CENTER)
    add_text(s4, x + Inches(0.1), y + Inches(1.3), Inches(1.5), Inches(0.5),
             sw, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s4, x + Inches(0.1), y + Inches(1.75), Inches(1.5), Inches(0.3),
             en, size=8, color=GOLD_LIGHT, align=PP_ALIGN.CENTER)
    add_text(s4, x + Inches(0.15), y + Inches(2.15), Inches(1.45), Inches(1.2),
             desc, size=8, color=INK_300, align=PP_ALIGN.CENTER)
    # Arrow between steps (except last)
    if i < len(steps) - 1:
        arr = s4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + Inches(1.74), y + Inches(1.5),
                                  Inches(0.2), Inches(0.3))
        fill(arr, GOLD)

# Bottom note
note = add_rect(s4, Inches(0.6), Inches(6.1), Inches(8.8), Inches(0.7), INK_800, rounded=True)
note.adjustments[0] = 0.3
add_text(s4, Inches(0.8), Inches(6.2), Inches(8.4), Inches(0.5),
         "💡 Wakati wote, mteja anaona unalindwa, na wewe unaona umelindwa.  /  At every step, both sides see protection.",
         size=11, color=GOLD_LIGHT, align=PP_ALIGN.CENTER)
footer(s4, 4, TOTAL_SLIDES)


# ── Slide 5 — Faida zako (Your benefits) ───────────────────────────────
s5 = prs.slides.add_slide(blank)
slide_bg(s5, INK_900)
add_pill(s5, Inches(0.6), Inches(0.5), Inches(1.4), Inches(0.4),
         "FAIDA ZAKO  /  YOUR BENEFITS", fill_rgb=EMERALD, text_rgb=WHITE, size=9)
add_text(s5, Inches(0.6), Inches(1.0), Inches(9), Inches(1.0),
         "Kwa nini Biz-Salama?",
         size=38, bold=True, color=WHITE)
add_text(s5, Inches(0.6), Inches(1.95), Inches(9), Inches(0.4),
         "Six concrete reasons to start selling on Biz-Salama today",
         size=12, color=INK_400)

# 3 × 2 grid of benefit cards
items = [
    ("🛡", "Hakuna chargebacks",     "No chargebacks",      "Mteja anapomaliza, pesa ni zako milele."),
    ("📲", "Malipo ya mara moja",   "Instant M-Pesa",      "Hakuna kungoja siku 3-5 za benki."),
    ("🎙", "Orodhesha kwa sauti",   "Voice listing",       "Sema bidhaa yako — tutaitafsiri na kuipublish."),
    ("👩🏾", "Walengwa wanaokujua",  "Verified buyers",     "Tunathibitisha namba ya simu ya kila mnunuzi."),
    ("⚖️", "Tukio la mgogoro?",     "Dispute resolution",  "Tunakaa katikati — uamuzi wa haki kwa pande zote."),
    ("🌍", "Toka Dar mpaka Mwanza", "Reach all of TZ",     "Tanzania nzima inakuona — siyo mtaa wako tu."),
]
for i, (emoji, sw, en, body) in enumerate(items):
    col, row = i % 3, i // 3
    x = Inches(0.6 + col * 3.0)
    y = Inches(2.6 + row * 2.0)
    card = add_rect(s5, x, y, Inches(2.85), Inches(1.85), INK_800, rounded=True)
    card.adjustments[0] = 0.1
    add_text(s5, x + Inches(0.2), y + Inches(0.15), Inches(0.6), Inches(0.5),
             emoji, size=22)
    add_text(s5, x + Inches(0.95), y + Inches(0.15), Inches(1.85), Inches(0.5),
             sw, size=12, bold=True, color=WHITE)
    add_text(s5, x + Inches(0.95), y + Inches(0.55), Inches(1.85), Inches(0.3),
             en, size=8, color=GOLD_LIGHT)
    add_text(s5, x + Inches(0.2), y + Inches(1.0), Inches(2.5), Inches(0.85),
             body, size=9, color=INK_300)
footer(s5, 5, TOTAL_SLIDES)


# ── Slide 6 — Ada / Fees ───────────────────────────────────────────────
s6 = prs.slides.add_slide(blank)
slide_bg(s6, INK_900)
add_pill(s6, Inches(0.6), Inches(0.5), Inches(1.0), Inches(0.4),
         "ADA  /  FEES", fill_rgb=GOLD, text_rgb=INK_900, size=9)
add_text(s6, Inches(0.6), Inches(1.0), Inches(9), Inches(1.0),
         "Ada wazi — hakuna siri",
         size=38, bold=True, color=WHITE)
add_text(s6, Inches(0.6), Inches(1.95), Inches(9), Inches(0.4),
         "Transparent fees — what you see is what you pay.",
         size=12, color=INK_400)

# Example transaction breakdown — based on real ledger numbers for TZS 100,000
add_text(s6, Inches(0.6), Inches(2.6), Inches(9), Inches(0.5),
         "Mfano  /  Example: bidhaa yako ya TZS 100,000",
         size=16, bold=True, color=GOLD_LIGHT)

bars = [
    ("Bei yako (asking price)", "Your price", "TZS 100,000", 1.0, WHITE),
    ("− Ada ya supply 2%", "− 2% supply fee", "− TZS 2,000", 0.02, INK_400),
    ("= Unachopata", "= You receive", "TZS 98,000 (98%)", 0.98, EMERALD_LT),
]
y = Inches(3.4)
for sw, en, amt, frac, col in bars:
    add_text(s6, Inches(0.6), y, Inches(4), Inches(0.4), sw, size=13, bold=True, color=col)
    add_text(s6, Inches(0.6), y + Inches(0.3), Inches(4), Inches(0.3), en, size=9, color=INK_400)
    # Bar
    bg = add_rect(s6, Inches(4.7), y + Inches(0.1), Inches(4.0), Inches(0.4), INK_800, rounded=True)
    bg.adjustments[0] = 0.5
    if frac > 0:
        fg = add_rect(s6, Inches(4.7), y + Inches(0.1), Inches(4.0 * frac), Inches(0.4), col, rounded=True)
        fg.adjustments[0] = 0.5
    add_text(s6, Inches(7.0), y + Inches(0.1), Inches(2.5), Inches(0.4),
             amt, size=11, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
    y += Inches(0.75)

# Side note: buyer pays the 3%
note = add_rect(s6, Inches(0.6), Inches(6.0), Inches(8.8), Inches(0.7), EMERALD, rounded=True)
note.adjustments[0] = 0.2
note.fill.transparency = 0.85
add_text(s6, Inches(0.8), Inches(6.1), Inches(8.4), Inches(0.5),
         "✅  Mteja anaongeza 3% juu ya bei yako kama ada ya escrow. Wewe haupotezi kitu nyongeza.  /  Buyer pays extra 3% — never deducted from your price.",
         size=11, color=WHITE, align=PP_ALIGN.CENTER)
footer(s6, 6, TOTAL_SLIDES)


# ── Slide 7 — 3-Party Escrow ───────────────────────────────────────────
s7 = prs.slides.add_slide(blank)
slide_bg(s7, INK_900)
add_pill(s7, Inches(0.6), Inches(0.5), Inches(1.4), Inches(0.4),
         "3-PARTY ESCROW  /  HAWKER MODE", fill_rgb=GOLD, text_rgb=INK_900, size=9)
add_text(s7, Inches(0.6), Inches(1.0), Inches(9), Inches(1.0),
         "Mchuuzi? Tunakulinda mara mbili.",
         size=34, bold=True, color=WHITE)
add_text(s7, Inches(0.6), Inches(1.95), Inches(9), Inches(0.4),
         "Hawker? You're protected on BOTH sides — supplier AND buyer.",
         size=13, color=INK_400)

# Three actor cards horizontally
actors = [
    ("👜", "MCHUUZI", "HAWKER (you)",       "Unaunganisha mteja na duka. Unapata commission ya 8-15%."),
    ("🏪", "MMILIKI WA DUKA", "SUPPLIER",   "Anauza bidhaa yake. Anapata bei kamili (-2%)."),
    ("🛍", "MNUNUZI", "BUYER",              "Analipa Biz-Salama. Anaona Letter of Comfort iliyosainiwa."),
]
for i, (emoji, sw, en, body) in enumerate(actors):
    x = Inches(0.6 + i * 3.0)
    y = Inches(2.6)
    card = add_rect(s7, x, y, Inches(2.85), Inches(2.7), INK_800, rounded=True)
    card.adjustments[0] = 0.1
    add_text(s7, x, y + Inches(0.3), Inches(2.85), Inches(0.7), emoji, size=44, align=PP_ALIGN.CENTER)
    add_text(s7, x, y + Inches(1.2), Inches(2.85), Inches(0.4), sw,
             size=13, bold=True, color=GOLD_LIGHT, align=PP_ALIGN.CENTER)
    add_text(s7, x, y + Inches(1.55), Inches(2.85), Inches(0.3), en,
             size=9, color=INK_400, align=PP_ALIGN.CENTER)
    add_text(s7, x + Inches(0.2), y + Inches(1.95), Inches(2.45), Inches(0.75),
             body, size=9, color=INK_300, align=PP_ALIGN.CENTER)

# Bottom message
note = add_rect(s7, Inches(0.6), Inches(5.6), Inches(8.8), Inches(1.2), GOLD, rounded=True)
note.adjustments[0] = 0.1
note.fill.transparency = 0.88
add_text(s7, Inches(0.8), Inches(5.75), Inches(8.4), Inches(0.5),
         "🔐  Letter of Comfort iliyosainiwa kidijiti",
         size=14, bold=True, color=GOLD_LIGHT, align=PP_ALIGN.CENTER)
add_text(s7, Inches(0.8), Inches(6.15), Inches(8.4), Inches(0.5),
         "Kila upande anaona maelezo yake tu — mmiliki haoni bei mteja anayolipa. Mteja anaona pesa zimeshikiliwa salama.",
         size=10, color=WHITE, align=PP_ALIGN.CENTER)
footer(s7, 7, TOTAL_SLIDES)


# ── Slide 8 — Voice Listing ────────────────────────────────────────────
s8 = prs.slides.add_slide(blank)
slide_bg(s8, INK_900)
add_pill(s8, Inches(0.6), Inches(0.5), Inches(1.4), Inches(0.4),
         "ORODHESHA KWA SAUTI  /  VOICE LISTING", fill_rgb=EMERALD, text_rgb=WHITE, size=9)
add_text(s8, Inches(0.6), Inches(1.0), Inches(9), Inches(1.0),
         "Sekunde 20 — bidhaa iko hewani",
         size=36, bold=True, color=WHITE)
add_text(s8, Inches(0.6), Inches(1.95), Inches(9), Inches(0.4),
         "20 seconds from idea to live listing — speak Swahili or English.",
         size=13, color=INK_400)

# Big mic icon
mic = s8.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.5), Inches(2.6), Inches(2), Inches(2))
fill(mic, GOLD)
add_text(s8, Inches(7.5), Inches(3.0), Inches(2), Inches(1.5),
         "🎙", size=90, align=PP_ALIGN.CENTER)

# 3 transcribed examples
examples = [
    "🗣 'Kitenge cha bei rahisi, mita 5, Kariakoo, TSh 12,000.'",
    "→  ✅  Kitenge fabric (5m), Kariakoo, TZS 12,000  (auto-categorized)",
    "",
    "🗣 'Samsung A54 mpya, warranty mwaka mzima, TSh 850,000.'",
    "→  ✅  Samsung Galaxy A54, electronics, TZS 850,000  (price + category)",
]
y = Inches(2.7)
for line in examples:
    if not line:
        y += Inches(0.25)
        continue
    is_result = line.startswith("→")
    add_text(s8, Inches(0.6), y, Inches(6.7), Inches(0.4),
             line, size=11, bold=is_result,
             color=EMERALD_LT if is_result else WHITE)
    y += Inches(0.4)

# Bottom note
note = add_rect(s8, Inches(0.6), Inches(6.0), Inches(8.8), Inches(0.7), INK_800, rounded=True)
note.adjustments[0] = 0.3
add_text(s8, Inches(0.8), Inches(6.1), Inches(8.4), Inches(0.5),
         "💡 Powered by OpenAI Whisper — inaelewa Kiswahili sanifu, lugha za mitaani, na lafudhi za Kanda.",
         size=11, color=GOLD_LIGHT, align=PP_ALIGN.CENTER)
footer(s8, 8, TOTAL_SLIDES)


# ── Slide 9 — Trust & Safety ───────────────────────────────────────────
s9 = prs.slides.add_slide(blank)
slide_bg(s9, INK_900)
add_pill(s9, Inches(0.6), Inches(0.5), Inches(1.4), Inches(0.4),
         "USALAMA  /  TRUST & SAFETY", fill_rgb=GOLD, text_rgb=INK_900, size=9)
add_text(s9, Inches(0.6), Inches(1.0), Inches(9), Inches(1.0),
         "Sababu watu wanakuamini",
         size=38, bold=True, color=WHITE)
add_text(s9, Inches(0.6), Inches(1.95), Inches(9), Inches(0.4),
         "Why buyers trust sellers on Biz-Salama (you get a verified badge).",
         size=12, color=INK_400)

trust = [
    ("🆔", "KYC iliyokamilika", "Verified KYC",
     "NIDA + TIN + Leseni ya biashara. Mara hii inathibitishwa, unapata 'Verified Seller' badge."),
    ("📜", "Cheti cha Escrow", "Escrow license",
     "Tunafanya kazi chini ya sheria za Tanzania (TRA + BoT compliant) — siyo kampuni ya nje."),
    ("🔍", "Ufuatiliaji wa udanganyifu", "Auto fraud detection",
     "Sheria 5 zinazochunguza kila order: velocity, self-deal, watchlist, refund-rate, new-account."),
    ("⚖️", "Mgogoro? Uamuzi wa siku 7", "7-day dispute auto-resolve",
     "Kama hakuna anayejibu, mfumo unatoa uamuzi wa haki — siyo lazima ujue mawakili."),
]
for i, (emoji, sw, en, body) in enumerate(trust):
    col, row = i % 2, i // 2
    x = Inches(0.6 + col * 4.6)
    y = Inches(2.6 + row * 1.95)
    card = add_rect(s9, x, y, Inches(4.3), Inches(1.75), INK_800, rounded=True)
    card.adjustments[0] = 0.1
    add_text(s9, x + Inches(0.25), y + Inches(0.2), Inches(0.6), Inches(0.5),
             emoji, size=24)
    add_text(s9, x + Inches(0.95), y + Inches(0.2), Inches(3.1), Inches(0.4),
             sw, size=13, bold=True, color=WHITE)
    add_text(s9, x + Inches(0.95), y + Inches(0.6), Inches(3.1), Inches(0.3),
             en, size=9, color=GOLD_LIGHT)
    add_text(s9, x + Inches(0.25), y + Inches(1.0), Inches(3.85), Inches(0.7),
             body, size=9, color=INK_300)
footer(s9, 9, TOTAL_SLIDES)


# ── Slide 10 — Mafanikio (Sample story) ────────────────────────────────
s10 = prs.slides.add_slide(blank)
slide_bg(s10, INK_900)
add_pill(s10, Inches(0.6), Inches(0.5), Inches(1.4), Inches(0.4),
         "MAFANIKIO  /  SUCCESS STORY", fill_rgb=EMERALD, text_rgb=WHITE, size=9)
add_text(s10, Inches(0.6), Inches(1.0), Inches(9), Inches(1.0),
         "Mama Asha — Kitenge Kariakoo",
         size=34, bold=True, color=WHITE)
add_text(s10, Inches(0.6), Inches(1.95), Inches(9), Inches(0.4),
         "How a Kariakoo fabric seller went from 3 buyers/month to 40+",
         size=12, color=INK_400)

# Big quote card
quote = add_rect(s10, Inches(0.6), Inches(2.6), Inches(5.2), Inches(3.6), INK_800, rounded=True)
quote.adjustments[0] = 0.08
add_text(s10, Inches(0.85), Inches(2.8), Inches(0.6), Inches(0.6),
         "❝", size=48, color=GOLD)
add_text(s10, Inches(1.4), Inches(2.85), Inches(4.2), Inches(2.7),
         "Hapo awali, wateja walikuwa wananiomba kuwatumia bidhaa kabla ya kulipa. Mara nyingi nilipotea pesa.\n\n"
         "Tangu nilipoanza Biz-Salama, mteja analipa kwanza — pesa zinashikiliwa, nikituma na kuthibitishwa, "
         "nalipwa kwa M-Pesa mara moja. Sasa nauza Kanda za Pwani, Arusha, mpaka Mbeya.",
         size=11, color=INK_300, italic=True)
add_text(s10, Inches(1.4), Inches(5.55), Inches(4.2), Inches(0.4),
         "— Asha M.,  Kariakoo, Dar es Salaam",
         size=10, bold=True, color=GOLD_LIGHT)

# Stats sidebar
stats = [
    ("3 → 40+",   "wateja kwa mwezi  /  buyers per month"),
    ("0%",        "chargebacks tangu kuanza  /  since joining"),
    ("< 24 hr",   "muda wa wastani wa malipo  /  avg payout time"),
    ("4.9 ★",     "rating ya wateja  /  buyer rating"),
]
y = Inches(2.6)
for big, small in stats:
    card = add_rect(s10, Inches(6.0), y, Inches(3.4), Inches(0.8), INK_800, rounded=True)
    card.adjustments[0] = 0.15
    add_text(s10, Inches(6.2), y + Inches(0.05), Inches(1.5), Inches(0.7),
             big, size=22, bold=True, color=GOLD_LIGHT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s10, Inches(7.7), y + Inches(0.05), Inches(1.7), Inches(0.7),
             small, size=8, color=INK_300, anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.9)
footer(s10, 10, TOTAL_SLIDES)


# ── Slide 11 — Anza sasa (Call to action) ─────────────────────────────
s11 = prs.slides.add_slide(blank)
slide_bg(s11, INK_900)
# Decorative
blob = s11.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7), Inches(-1.5), Inches(5), Inches(5))
fill(blob, GOLD)
blob.fill.transparency = 0.85

add_pill(s11, Inches(0.6), Inches(0.5), Inches(1.0), Inches(0.4),
         "ANZA SASA  /  GET STARTED", fill_rgb=EMERALD, text_rgb=WHITE, size=9)
add_text(s11, Inches(0.6), Inches(1.1), Inches(9), Inches(1.2),
         "Anza biashara salama leo",
         size=52, bold=True, color=WHITE)
add_text(s11, Inches(0.6), Inches(2.3), Inches(9), Inches(0.5),
         "Start selling safely today — onboarding takes ~10 minutes.",
         size=16, color=GOLD_LIGHT)

# 4-step onboarding checklist
steps = [
    ("1", "Tembelea  /  Visit",        "biz-salama.co.tz/register"),
    ("2", "Sajili akaunti  /  Sign up", "Namba ya simu + nenosiri (10 sec)"),
    ("3", "Pakizia hati  /  Upload docs", "NIDA, TIN, leseni — picha 5 kwa simu"),
    ("4", "Anza kuuza  /  Start selling", "Orodhesha bidhaa ya kwanza — kwa sauti au kwa kuandika"),
]
y = Inches(3.3)
for n, sw_en, body in steps:
    card = add_rect(s11, Inches(0.6), y, Inches(6.5), Inches(0.6), INK_800, rounded=True)
    card.adjustments[0] = 0.2
    num = s11.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.75), y + Inches(0.1),
                                Inches(0.4), Inches(0.4))
    fill(num, GOLD)
    add_text(s11, Inches(0.75), y + Inches(0.12), Inches(0.4), Inches(0.4),
             n, size=14, bold=True, color=INK_900, align=PP_ALIGN.CENTER)
    add_text(s11, Inches(1.35), y + Inches(0.05), Inches(2.4), Inches(0.5),
             sw_en, size=12, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s11, Inches(3.75), y + Inches(0.05), Inches(3.3), Inches(0.5),
             body, size=10, color=GOLD_LIGHT, anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.75)

# Big CTA card on right
cta = add_rect(s11, Inches(7.4), Inches(3.3), Inches(2.0), Inches(2.7), GOLD, rounded=True)
cta.adjustments[0] = 0.1
add_text(s11, Inches(7.4), Inches(3.5), Inches(2.0), Inches(0.4),
         "🚀", size=36, align=PP_ALIGN.CENTER)
add_text(s11, Inches(7.4), Inches(4.2), Inches(2.0), Inches(0.4),
         "ANZA",
         size=20, bold=True, color=INK_900, align=PP_ALIGN.CENTER)
add_text(s11, Inches(7.4), Inches(4.6), Inches(2.0), Inches(0.4),
         "START NOW",
         size=11, bold=True, color=INK_800, align=PP_ALIGN.CENTER)
add_text(s11, Inches(7.4), Inches(5.1), Inches(2.0), Inches(0.6),
         "biz-salama.co.tz\n/onboard/seller",
         size=9, color=INK_900, align=PP_ALIGN.CENTER)

# Contact strip
contact = add_rect(s11, Inches(0.6), Inches(6.5), Inches(8.8), Inches(0.5), INK_800, rounded=True)
contact.adjustments[0] = 0.3
add_text(s11, Inches(0.8), Inches(6.58), Inches(8.4), Inches(0.4),
         "📞  +255 700 123 456    ·    ✉️  support@biz-salama.co.tz    ·    🌐  biz-salama.co.tz",
         size=10, color=GOLD_LIGHT, align=PP_ALIGN.CENTER)
footer(s11, 11, TOTAL_SLIDES)


# ─── Save ────────────────────────────────────────────────────────────────
out = Path("/app/biz_salama_seller_pitch_sw.pptx")
prs.save(out)
print(f"✅ Pitch deck saved to {out}")
print(f"   Slides: {TOTAL_SLIDES}  ·  Size: {out.stat().st_size // 1024} KB")
